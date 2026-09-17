"""NeuroRedTeam Phase-1 deck, built on the RVCE EL template (template3 = original + 2 results slides + 2nd refs slide).
Plain wording, no em dashes. Output: C:/Users/hp/Desktop/ConnectomeLab/NeuroRedTeam_Phase1.pptx"""
import copy, os
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
SRC = os.path.join(HERE, "template3.pptx")
OUT = os.path.join(ROOT, "NeuroRedTeam_Phase1.pptx")
NAVY, RED = RGBColor(0x00, 0x20, 0x60), RGBColor(0xC0, 0x00, 0x00)

prs = Presentation(SRC); S = list(prs.slides)

def body_ph(slide):
    for sh in slide.shapes:
        if sh.is_placeholder and sh.placeholder_format.type == 2: return sh
    raise KeyError
def title_box(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and not sh.is_placeholder and sh.top < Inches(1.3): return sh
    raise KeyError
_b4 = body_ph(S[3])._element
PPR_L0 = copy.deepcopy(_b4.findall(".//a:p", NS)[1].find("a:pPr", NS))
PPR_L1 = copy.deepcopy(body_ph(S[12])._element.findall(".//a:p", NS)[2].find("a:pPr", NS))
PPR_NONE = copy.deepcopy(_b4.findall(".//a:p", NS)[2].find("a:pPr", NS))

def make_para(level, runs, size, first=False):
    p = etree.Element("{%s}p" % NS["a"]); ppr = copy.deepcopy({0: PPR_L0, 1: PPR_L1, "none": PPR_NONE}[level])
    if first: ppr.find("a:spcBef", NS).find("a:spcPts", NS).set("val", "0")
    p.append(ppr)
    for text, bold in runs:
        r = etree.SubElement(p, "{%s}r" % NS["a"]); rpr = etree.SubElement(r, "{%s}rPr" % NS["a"], lang="en-IN", sz=str(size * 100))
        if bold: rpr.set("b", "1")
        t = etree.SubElement(r, "{%s}t" % NS["a"]); t.text = text
        if text != text.strip(): t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    etree.SubElement(p, "{%s}endParaRPr" % NS["a"], sz=str(size * 100)); return p

def set_body(slide, items, size=18, sub_size=None):
    sub_size = sub_size or size - 2
    txBody = body_ph(slide)._element.find(".//p:txBody", NS)
    for p in txBody.findall("a:p", NS): txBody.remove(p)
    bp = txBody.find("a:bodyPr", NS)
    for c in list(bp): bp.remove(c)
    etree.SubElement(bp, "{%s}noAutofit" % NS["a"])
    for i, (level, runs) in enumerate(items):
        if isinstance(runs, str): runs = [(runs, False)]
        txBody.append(make_para(level, runs, size if level in (0, "none") else sub_size, first=(i == 0)))

def set_simple_text(shape, text, size=None):
    txBody = shape._element.find(".//p:txBody", NS); paras = txBody.findall("a:p", NS); p0 = paras[0]
    runs = p0.findall("a:r", NS); rpr = copy.deepcopy(runs[0].find("a:rPr", NS))
    for extra in paras[1:]: txBody.remove(extra)
    for r in runs: p0.remove(r)
    for li, line in enumerate(text.split("\n")):
        target = p0 if li == 0 else copy.deepcopy(p0)
        if li > 0:
            for r in target.findall("a:r", NS): target.remove(r)
            txBody.append(target)
        r = etree.Element("{%s}r" % NS["a"]); rp = copy.deepcopy(rpr)
        if size: rp.set("sz", str(size * 100))
        r.append(rp); t = etree.SubElement(r, "{%s}t" % NS["a"]); t.text = line
        e = target.find("a:endParaRPr", NS)
        target.insert(list(target).index(e), r) if e is not None else target.append(r)

def set_slide_number(slide, n):
    for sh in slide.shapes:
        if sh.is_placeholder and sh.placeholder_format.type == 13:
            for r in sh._element.findall(".//a:r", NS): r.find("a:t", NS).text = ""
            for f in sh._element.findall(".//a:fld", NS): f.find("a:t", NS).text = str(n)

def clear_body(slide):
    txBody = body_ph(slide)._element.find(".//p:txBody", NS)
    for p in txBody.findall("a:p", NS): txBody.remove(p)
    txBody.append(make_para("none", [("", False)], 10))

def caption(slide, x, y, w, h, text, size=13, color=NAVY):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text; r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = "Bookman Old Style"

# ============ 1. Title ============
s = S[0]
boxes = {sh.text_frame.text.strip(): sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip() in ("Title", "Theme")}
set_simple_text(boxes["Title"], "NeuroRedTeam: Attack and Repair Testing of a Real Fly Brain\nHow fragile is a brain's behaviour, and how do you fix it?", size=24)
set_simple_text(boxes["Theme"], "Theme: Artificial Intelligence, Neural Networks and Computational Neuroscience\nTeam: Govind  |  Bharath  |  Dhanush A  |  Shreenidhi Chalmi", size=20)
for b, sz in ((boxes["Title"], "2000"), (boxes["Theme"], "1800")):
    for r in b._element.findall(".//a:p", NS)[1].findall("a:r", NS):
        rp = r.find("a:rPr", NS); rp.set("sz", sz)
        sf = rp.find("a:solidFill", NS)
        if sf is not None: rp.remove(sf)
        sf = etree.Element("{%s}solidFill" % NS["a"]); etree.SubElement(sf, "{%s}srgbClr" % NS["a"], val="FFFFFF"); rp.insert(0, sf)

# ============ 4. Introduction ============
set_body(S[3], [
    (0, [("The starting point: ", True),
         ("In October 2024 scientists finished mapping every neuron and every connection in a fruit fly's brain: 139,255 neurons, "
          "15 million connections, released free. It is the first complete wiring diagram of an animal that can walk, fly, see and learn. "
          "A follow-up paper showed the whole map can be run on a normal laptop and it behaves like a real fly.", False)]),
    (0, [("What we already have running: ", True),
         ("we reproduced that simulation on our laptop. Give the simulated fly sugar and its feeding neuron fires; give it random input "
          "and nothing happens. Cut 3 specific neurons out of 139,255 and feeding stops; cut 10 random ones and nothing changes (slides 10 and 11).", False)]),
    (0, [("The question we ask: ", True),
         ("that 3-neuron result was found by hand, for one behaviour. Nobody knows, for the brain as a whole, how few targeted changes it takes "
          "to break a behaviour, or how few it takes to repair it. We treat the brain like a system under security testing: one AI attacks, "
          "one AI repairs, and we record the results for every behaviour we can measure.", False)]),
    (0, [("Why it matters: ", True),
         ("it tells neuroscientists which circuits are fragile and which are protected, it is a first step toward designing repair stimulation "
          "(the logic behind deep brain stimulation) on a complete brain, and every result is a prediction a real fly lab can test.", False)]),
], size=17)

# ============ 5. Literature Review ============
set_body(S[4], [
    (0, [("The map and the simulation:", True)]),
    (1, [("Dorkenwald et al. (2024, Nature): ", True), ("the complete fly brain connectome (FlyWire).", False)]),
    (1, [("Shiu et al. (2024, Nature): ", True), ("ran the whole map with simple spiking neurons on a laptop and correctly predicted feeding, grooming and more. We reproduced this.", False)]),
    (0, [("Perturbing the brain model (the closest prior work):", True)]),
    (1, [("Shiu et al. (2024) and a bioRxiv preprint (Dec 2025): ", True), ("silenced hand-picked neurons in the same model to test specific ideas; the 2025 work ran combinatorial silencing of two inhibitors and found feeding suppression is redundant. Manual, one circuit, no search for minimal sets, no repair.", False)]),
    (1, [("Pospisil et al. (2024, Nature, the effectome): ", True), ("a linear model of the whole fly brain that predicts what perturbing one neuron does to others.", False)]),
    (1, [("Yan et al. (2017, Nature): ", True), ("control theory on the 302-neuron worm to find the minimal driver neurons for movement; one prediction was confirmed by laser ablation in real worms. Worm only, linear maths, no behaviour simulation, no repair.", False)]),
    (0, [("AI agents that run experiments: ", True), ("exist for chemistry, materials science and quantum physics (2025 to 2026). For the fly brain only single-command tools exist (type \"give sugar\", get one run). No agent chooses its own experiments on a brain.", False)]),
    (0, [("Gap: ", True), ("nobody has automatically searched a whole-brain behavioural simulation for the smallest change that breaks a behaviour, nobody has designed the smallest stimulation that repairs it, and there is no table of how fragile each fly behaviour is.", False)]),
], size=14, sub_size=13)

# ============ 6. Problem Definition ============
set_body(S[5], [
    (0, [("The problem: ", True),
         ("we have every wire in a brain, but we do not know how fragile its behaviours are. Prior work silenced a few hand-picked neurons at a time. "
          "There is no systematic answer to: how few targeted changes can hijack a behaviour, which neurons are the weak points, and what is the smallest "
          "intervention that restores normal function.", False)]),
    (0, [("Why it matters if left unanswered: ", True),
         ("neuroscientists cannot tell which circuits are protected by redundancy and which are single points of failure; nobody can design repair "
          "stimulation on a complete brain; and the mouse and human brain maps now being built will hit the same gap with far larger networks.", False)]),
    (0, [("Problem statement: ", True),
         ("Build a laptop platform that runs the complete fruit-fly connectome as a spiking neural network with measurable behaviours (eat, refuse bitter, "
          "groom, escape). On it, build an attacker agent that finds the minimal constrained set of perturbations that forces or blocks a behaviour, and a "
          "defender agent that finds the minimal stimulation that restores it. Produce a fragility atlas of fly behaviours and a benchmark comparing "
          "random search, greedy search and LLM agents on this task.", False)]),
], size=16)

# ============ 7. Objectives ============
set_body(S[6], [
    (0, [("Main objective: ", True),
         ("measure, for a complete real brain, how easy each behaviour is to break and how easy it is to repair, using two AI agents that play attacker and defender.", False)]),
    (0, [("Specific objectives:", True)]),
    (1, [("Platform: ", True), ("fast simulator of all 139,255 neurons (stimulate, record, cut, weaken, inject), a name table for neurons, and four measurable behaviours reproduced from the 2024 Nature paper.", False)]),
    (1, [("Attack rules: ", True), ("a written attack surface so the problem is not trivial: no touching motor neurons, only silence, weaken or inject at realistic rates, with a budget.", False)]),
    (1, [("Attacker agent: ", True), ("finds the smallest set of allowed changes that forces a behaviour (eat with no sugar), blocks it (no eating with sugar) or corrupts it (eat bitter). Three versions: random, greedy, LLM with tools.", False)]),
    (1, [("Defender agent: ", True), ("given the attacked brain, finds the smallest set of neurons to stimulate that restores the behaviour, and proves it did not simply force it (must work with sugar and stay quiet without).", False)]),
    (1, [("Fragility atlas and benchmark: ", True), ("for each behaviour: minimal attack size, minimal repair size, key neurons, redundancy; plus how many simulations each method needed.", False)]),
    (1, [("Virtual fly and dashboard: ", True), ("a 2-D fly driven by the brain and one web page: press Attack, watch neurons turn red and the fly stop eating; press Defend, watch it eat again.", False)]),
    (0, [("What is new: ", True), ("the automated minimal attack search, the defender, and the atlas. The simulator and virtual fly are our platform, reproduced from published work, not claimed as new.", False)]),
], size=15, sub_size=14)

# ============ 8. Methodology ============
set_body(S[7], [
    (0, [("Design: ", True), ("simulation experiments with controls. Every attack and repair is repeated over several random seeds; results are checked against the neurons named in published papers.", False)]),
    (0, [("Data: ", True), ("public FlyWire v783 connectome (neurons, synapse counts, excitatory or inhibitory sign, cell-type labels). No data collected from people or animals.", False)]),
    (0, [("Behaviours as numbers: ", True), ("eat = feeding neuron MN9 above 60 Hz with sugar; refuse = MN9 below 15 Hz with sugar plus bitter; groom and escape read from their descending neurons. Each is a pass or fail test the agents must satisfy.", False)]),
    (0, [("How an agent works: ", True), ("it is a decision loop with tools: look up who connects to whom, run the brain with chosen changes, read the result, choose the next experiment, stop at the goal. The rules are enforced in code, so the agent cannot cheat. Random and greedy versions use the same tools without an LLM, so we can measure whether the LLM is actually better.", False)]),
    (0, [("Speed: ", True), ("300 ms runs on the sub-network reachable from the stimulus, about 1 to 2 s each, with caching. An agent run of 30 experiments takes about a minute; the full atlas runs overnight.", False)]),
    (0, [("Tools: ", True), ("Python, NumPy and SciPy sparse, NetworkX, Claude or OpenAI API for the agents, pygame for the fly, FastAPI and a web page for the dashboard. Laptop only.", False)]),
    (0, [("Team: ", True), ("simulator speed-ups and virtual fly (Govind); name table, behaviours and validation (Bharath); attacker agent (Dhanush); attack rules, defender, benchmark scoring and dashboard (Shreenidhi).", False)]),
], size=14)

# ============ 9. Flowchart ============
s = S[8]; clear_body(s)
def box(slide, x, y, w, h, title, sub, fill=NAVY):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill; shp.line.color.rgb = RGBColor(0x42, 0x71, 0x9B); shp.line.width = Pt(1.25); shp.shadow.inherit = False
    tf = shp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08); tf.margin_top = tf.margin_bottom = Inches(0.05)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title; r.font.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(255, 255, 255); r.font.name = "Bookman Old Style"
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = sub; r2.font.size = Pt(11); r2.font.color.rgb = RGBColor(255, 255, 255); r2.font.name = "Bookman Old Style"
def arrow(slide, x1, y1, x2, y2, head=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); c.line.color.rgb = RED; c.line.width = Pt(2.25)
    if head: etree.SubElement(c.line._get_or_add_ln(), "{%s}tailEnd" % NS["a"], type="triangle", w="med", len="med")
row1 = [("1. Load the brain", "FlyWire data: 139,255 neurons,\n15 M connections   (done)"),
        ("2. Simulator", "spiking neurons; stimulate,\nrecord, cut, weaken, inject   (done)"),
        ("3. Behaviours", "eat (done), refuse bitter,\ngroom, escape + name table"),
        ("4. Attack rules", "what the attacker may touch,\nbudget, realistic rates")]
row2 = [("5. Attacker agent", "smallest change that forces,\nblocks or corrupts a behaviour"),
        ("6. Defender agent", "smallest stimulation that\nrestores it, without forcing"),
        ("7. Fragility atlas", "behaviour x attack size x\nrepair size x key neurons"),
        ("8. Dashboard + paper", "Attack / Defend buttons,\nvirtual fly, report")]
Wb, Hb, GAP, X0, Y1, Y2 = 2.75, 1.25, 0.42, 0.55, 1.55, 4.15
for i, (t, sub) in enumerate(row1):
    x = X0 + i * (Wb + GAP); box(s, x, Y1, Wb, Hb, t, sub)
    if i < 3: arrow(s, x + Wb, Y1 + Hb / 2, x + Wb + GAP, Y1 + Hb / 2)
for i, (t, sub) in enumerate(row2):
    x = X0 + i * (Wb + GAP); box(s, x, Y2, Wb, Hb, t, sub, fill=RED if i < 3 else NAVY)
    if i < 3: arrow(s, x + Wb, Y2 + Hb / 2, x + Wb + GAP, Y2 + Hb / 2)
x4, x5, ymid = X0 + 3 * (Wb + GAP) + Wb / 2, X0 + Wb / 2, (Y1 + Hb + Y2) / 2
arrow(s, x4, Y1 + Hb, x4, ymid); arrow(s, x4, ymid, x5, ymid, head=False); arrow(s, x5, ymid, x5, Y2)
caption(s, 0.55, 5.65, 12.2, 0.9, "Red boxes are the new contributions. Steps 1 and 2 and the first behaviour are done (September 2026). "
        "Every attack and repair is repeated over several seeds and checked against published circuits.", size=14)

# ============ 10. Preliminary results ============
s = S[9]; set_simple_text(title_box(s), "Preliminary Results"); clear_body(s); set_slide_number(s, 10)
s.shapes.add_picture(os.path.join(ROOT, "results", "sugar_test.png"), Inches(0.45), Inches(1.35), width=Inches(6.6))
s.shapes.add_picture(os.path.join(ROOT, "results", "lesion_test.png"), Inches(7.35), Inches(1.35), width=Inches(5.6))
caption(s, 0.45, 4.05, 6.6, 0.6, "Sugar on the tongue: about 380 neurons active, feeding neuron MN9 at 95 to 98 Hz.\nRandom input: MN9 silent (0 Hz). Same result in 3 of 3 runs.", size=12)
caption(s, 7.35, 4.05, 5.6, 0.6, "First attack, done by hand: cut the 3 command neurons and feeding stops (1 Hz).\nCut 10 random neurons and nothing changes (91 Hz).", size=11)
set_body(s, [
    (0, [("What this shows: ", True), ("the real 139,255-neuron brain runs on a normal laptop and matches the Nature 2024 numbers; and the first entry of the fragility atlas already exists: feeding can be blocked with 3 neurons, but not with 10 random ones. The attacker agent's job is to find results like this automatically, for every behaviour, under rules.", False)]),
], size=15)
body_ph(s).top = Inches(4.85); body_ph(s).height = Inches(1.9)

# ============ 11. Live demo ============
s = S[10]; set_simple_text(title_box(s), "Live Demo"); clear_body(s); set_slide_number(s, 11)
s.shapes.add_picture(os.path.join(ROOT, "animations", "fly_feeding_demo.gif"), Inches(0.9), Inches(1.15), width=Inches(9.3))
set_body(s, [
    (0, [("Animation (plays in slideshow): ", True), ("left, the real wiring with neurons flashing as they fire, placed by distance from the tongue. Right, a fly whose mouth is driven live by the MN9 feeding neuron.", False)]),
    (0, [("0 to 350 ms: ", True), ("sugar arrives, the signal spreads hop by hop, MN9 fires, the mouth extends.", False)]),
    (0, [("350 ms, the attack: ", True), ("3 command neurons are cut (red X), MN9 goes silent, the mouth retracts.", False)]),
    (0, [("December: ", True), ("press Attack and the agent finds this itself; press Defend and it finds the neurons to stimulate so the fly eats again.", False)]),
], size=13, sub_size=12)
b = body_ph(s); b.left = Inches(10.3); b.top = Inches(1.15); b.width = Inches(2.85); b.height = Inches(5.5)

# ============ 12. Proposed Outcomes ============
set_body(S[11], [
    (0, [("Expected results:", True)]),
    (1, "A fragility atlas for four fly behaviours: how many targeted changes break each one, how many stimulated neurons repair it, which neurons keep appearing, and how redundant each circuit is."),
    (1, "An attacker and a defender agent, each in three versions (random, greedy, LLM with tools), with a benchmark table: intervention size and simulations used. This answers whether an LLM agent is actually better than a plain search."),
    (1, "Validation: the neurons our attacker finds compared with those named in the published papers (for feeding we can already check: the 3 command neurons)."),
    (1, "The platform: whole-brain simulator matching Nature 2024, a 2-D virtual fly, one dashboard with Attack and Defend buttons."),
    (0, [("Impact: ", True), ("a map of which brain circuits are protected and which are single points of failure; a first method for designing repair stimulation on a complete brain; testable predictions for fly labs; a testbed that transfers to the mouse and human maps being built now.", False)]),
    (0, [("Future work: ", True), ("more behaviours, the male brain (released June 2026) for a male vs female fragility comparison, and running attacks on a neuromorphic chip.", False)]),
], size=15, sub_size=14)

# ============ 13. Timeline ============
set_body(S[12], [
    (0, [("Phase 1: Foundation (Weeks 1 to 2, 17 Sep to 1 Oct)", True)]),
    (1, "Scope agreed with mentor; GitHub repo; simulator sped up (300 ms runs, sub-network mode, caching); neuron name table loaded."),
    (0, [("Phase 2: Behaviours and rules (Weeks 3 to 5, 1 Oct to 22 Oct)", True)]),
    (1, "Reproduce eat, refuse bitter, groom, escape as pass or fail tests; write the attack rules; greedy attacker as the first baseline."),
    (0, [("Phase 3: Attacker and defender agents (Weeks 6 to 9, 22 Oct to 19 Nov)", True)]),
    (1, "Simulator exposed as tools; LLM attacker with budget; defender with the two-condition repair test; random and greedy versions of both."),
    (0, [("Phase 4: Atlas, virtual fly, dashboard (Weeks 8 to 11, 5 Nov to 3 Dec)", True)]),
    (1, "Batch runs over behaviours, goals, methods and seeds; fragility atlas and benchmark table; 2-D fly; web page with Attack and Defend."),
    (0, [("Phase 5: Report and paper (Weeks 11 to 13, 3 Dec to 17 Dec)", True)]),
    (1, "Final demo, project report, open-source release, IEEE-format paper draft."),
], size=16, sub_size=15)

# ============ 14. Conclusion ============
set_body(S[13], [
    (0, [("Summary: ", True), ("the complete fly brain runs on our laptop and we have already broken one behaviour by hand with 3 neurons. NeuroRedTeam turns that into a method: an attacker agent that finds the smallest change that breaks a behaviour, a defender agent that finds the smallest stimulation that repairs it, and a fragility atlas of the whole brain. Prior work perturbed the model by hand; nobody has automated the attack, built the repair, or produced the atlas.", False)]),
    (0, [("Final thought: ", True), ("this is where neuroscience, AI and security thinking meet. Treating a brain like a system under test is new, and every answer it gives is a prediction a real lab can check.", False)]),
    (0, [("Next steps: ", True), ("mentor approval; behaviours and rules in October; agents in November; atlas, dashboard and paper in December.", False)]),
], size=18)

# ============ 15 to 16. References ============
REFS_A = [
    "Dorkenwald, S., Matsliah, A., Sterling, A. R., Schlegel, P., Yu, S.-C., McKellar, C. E., et al. (2024). Neuronal wiring diagram of an adult brain. Nature, 634, 124-138. https://doi.org/10.1038/s41586-024-07558-y",
    "Shiu, P. K., Sterne, G. R., Spiller, N., Blagburn, D. I., Mosca, T. J., et al. (2024). A Drosophila computational brain model reveals sensorimotor processing. Nature, 634, 210-219. https://doi.org/10.1038/s41586-024-07763-9",
    "Pospisil, D. A., Aragon, M. J., Dorkenwald, S., et al. (2024). The fly connectome reveals a path to the effectome. Nature, 634, 201-209. https://doi.org/10.1038/s41586-024-07982-0",
    "Schlegel, P., Yin, Y., Bates, A. S., et al. (2024). Whole-brain annotation and multi-connectome cell typing of Drosophila. Nature, 634, 139-152. https://doi.org/10.1038/s41586-024-07686-5",
    "Lin, A., Yang, R., Dorkenwald, S., et al. (2024). Network statistics of the whole-brain connectome of Drosophila. Nature, 634, 153-165. https://doi.org/10.1038/s41586-024-07968-y",
    "Yan, G., Vertes, P. E., Towlson, E. K., Chew, Y. L., Walker, D. S., Schafer, W. R., & Barabasi, A.-L. (2017). Network control principles predict neuron function in the Caenorhabditis elegans connectome. Nature, 550, 519-523. https://doi.org/10.1038/nature24056",
    "Whole-brain connectomics of Drosophila reveals a robust, distributed architecture for the suppression of feeding during escape. (2025). bioRxiv preprint. https://doi.org/10.64898/2025.12.14.694122",
    "Bates, A. S., Phelps, J. S., Kim, M., Yang, H. H., et al. (2026). Distributed control circuits across a brain-and-cord connectome. Nature. https://doi.org/10.1038/s41586-026-10735-w",
]
REFS_B = [
    "Lappalainen, J. K., Tschopp, F. D., Prakhya, S., et al. (2024). Connectome-constrained networks predict neural activity across the fly visual system. Nature, 634, 1132-1140. https://doi.org/10.1038/s41586-024-07939-3",
    "Vaxenburg, R., Siwanowicz, I., Merel, J., et al. (2025). Whole-body physics simulation of fruit fly locomotion. Nature, 643, 1312-1320. https://doi.org/10.1038/s41586-025-09029-4",
    "Suarez, L. E., Mihalik, A., Milisav, F., Marshall, K., Li, M., Vertes, P. E., Lajoie, G., & Misic, B. (2024). Connectome-based reservoir computing with the conn2res toolbox. Nature Communications, 15, 656. https://doi.org/10.1038/s41467-024-44900-4",
    "Jin, Z., Zhu, Y., Zhang, C., & Sui, Y. (2026). Whole-brain connectomic graph model enables whole-body locomotion control in fruit fly. arXiv:2602.17997.",
    "Dhiman, N. (2026). Topological sensitivity in connectome-constrained neural networks. arXiv:2604.04033.",
    "Neuromorphic simulation of Drosophila melanogaster brain connectome on Loihi 2. (2025). arXiv:2508.16792.",
    "FlyEM Project Team (HHMI Janelia, University of Cambridge, MRC LMB, Google Research). (2026). Sexual dimorphism in the complete connectome of the Drosophila male central nervous system. Cell. https://www.cell.com/cell/fulltext/S0092-8674(26)00942-6",
    "FlyWire Consortium. (2024). FlyWire Codex, release v783 [Dataset]. https://codex.flywire.ai",
    "Shiu, P. K. (2024). Drosophila_brain_model [Source code]. GitHub. https://github.com/philshiu/Drosophila_brain_model",
]
set_body(S[14], [("none", r) for r in REFS_A], size=11, sub_size=11); set_simple_text(title_box(S[14]), "References (1/2)")
set_body(S[15], [("none", r) for r in REFS_B], size=11, sub_size=11); set_simple_text(title_box(S[15]), "References (2/2)")

# ============ 17. Expected outcome ============
s = S[16]
set_body(s, [
    (0, [("Our target outcome: ", True), ("a conference publication plus an open-source release.", False)]),
    (0, [("Conference publication: ", True), ("an IEEE-format paper, \"NeuroRedTeam: automated attack and repair testing of a whole-brain connectome model\", reporting the fragility atlas and the agent benchmark. Target: IEEE conferences on computational intelligence or neural networks, or a computational neuroscience workshop.", False)]),
    (0, [("Open-source release: ", True), ("the simulator, the behaviour tests, the attack rules, both agents, the atlas and the dashboard on GitHub, so other students and labs can reproduce and extend it.", False)]),
    (0, [("Research proposal (stretch): ", True), ("extend the method to the male fly brain and to the mouse connectome, with a budget for a small GPU workstation.", False)]),
    (0, [("Patent: ", True), ("not planned. The value here is in open science and publication.", False)]),
], size=17)
for sh in s.shapes:
    if sh.is_placeholder and sh.placeholder_format.type == 16:
        for t in sh._element.findall(".//a:t", NS): t.text = "Friday, 18 September 2026"
for i, sl in enumerate(S, 1):
    if i >= 10: set_slide_number(sl, i)

prs.save(OUT); print("saved", OUT)
