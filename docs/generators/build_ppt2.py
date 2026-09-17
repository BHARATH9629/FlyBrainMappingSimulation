"""Fill the RVCE EL Phase-1 template (template3 = original + 2 results slides + 2nd references slide)
with the Connectome Lab content. Only text inside existing boxes is replaced; pictures/flowchart added."""
import copy
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
SRC, OUT = "template3.pptx", "EL_Phase1_Connectome_Lab.pptx"
NAVY, RED = RGBColor(0x00, 0x20, 0x60), RGBColor(0xC0, 0x00, 0x00)

prs = Presentation(SRC)
S = list(prs.slides)   # 0-based

def shapes_with_text(slide):
    return [sh for sh in slide.shapes if sh.has_text_frame]

def body_ph(slide):
    for sh in slide.shapes:
        if sh.is_placeholder and sh.placeholder_format.type == 2:
            return sh
    raise KeyError("no body")

def title_box(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and not sh.is_placeholder and sh.top < Inches(1.3):
            return sh
    raise KeyError("no title")

# paragraph property templates cloned from the original template
_b4, _b13 = body_ph(S[3])._element, body_ph(S[12])._element
PPR_L0 = copy.deepcopy(_b4.findall(".//a:p", NS)[1].find("a:pPr", NS))
PPR_L1 = copy.deepcopy(body_ph(S[12])._element.findall(".//a:p", NS)[2].find("a:pPr", NS))
PPR_NONE = copy.deepcopy(_b4.findall(".//a:p", NS)[2].find("a:pPr", NS))

def make_para(level, runs, size, first=False):
    src = {0: PPR_L0, 1: PPR_L1, "none": PPR_NONE}[level]
    p = etree.Element("{%s}p" % NS["a"])
    ppr = copy.deepcopy(src)
    if first:
        ppr.find("a:spcBef", NS).find("a:spcPts", NS).set("val", "0")
    p.append(ppr)
    for text, bold in runs:
        r = etree.SubElement(p, "{%s}r" % NS["a"])
        rpr = etree.SubElement(r, "{%s}rPr" % NS["a"], lang="en-IN", sz=str(size * 100))
        if bold:
            rpr.set("b", "1")
        t = etree.SubElement(r, "{%s}t" % NS["a"]); t.text = text
        if text != text.strip():
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    etree.SubElement(p, "{%s}endParaRPr" % NS["a"], sz=str(size * 100))
    return p

def set_body(slide, items, size=18, sub_size=None):
    sub_size = sub_size or size - 2
    txBody = body_ph(slide)._element.find(".//p:txBody", NS)
    for p in txBody.findall("a:p", NS):
        txBody.remove(p)
    bp = txBody.find("a:bodyPr", NS)
    for c in list(bp):
        bp.remove(c)
    etree.SubElement(bp, "{%s}noAutofit" % NS["a"])
    for i, (level, runs) in enumerate(items):
        if isinstance(runs, str):
            runs = [(runs, False)]
        txBody.append(make_para(level, runs, size if level in (0, "none") else sub_size, first=(i == 0)))

def set_simple_text(shape, text, size=None, color=None):
    txBody = shape._element.find(".//p:txBody", NS)
    paras = txBody.findall("a:p", NS); p0 = paras[0]
    runs = p0.findall("a:r", NS); rpr = copy.deepcopy(runs[0].find("a:rPr", NS))
    for extra in paras[1:]:
        txBody.remove(extra)
    for r in runs:
        p0.remove(r)
    for li, line in enumerate(text.split("\n")):
        target = p0 if li == 0 else copy.deepcopy(p0)
        if li > 0:
            for r in target.findall("a:r", NS):
                target.remove(r)
            txBody.append(target)
        r = etree.Element("{%s}r" % NS["a"]); rp = copy.deepcopy(rpr)
        if size: rp.set("sz", str(size * 100))
        if color:
            sf = rp.find("a:solidFill", NS)
            if sf is not None: rp.remove(sf)
            sf = etree.Element("{%s}solidFill" % NS["a"]); etree.SubElement(sf, "{%s}srgbClr" % NS["a"], val=color)
            rp.insert(0, sf)
        r.append(rp); t = etree.SubElement(r, "{%s}t" % NS["a"]); t.text = line
        e = target.find("a:endParaRPr", NS)
        target.insert(list(target).index(e), r) if e is not None else target.append(r)

def set_slide_number(slide, n):
    for sh in slide.shapes:
        if sh.is_placeholder and sh.placeholder_format.type == 13:
            for r in sh._element.findall(".//a:r", NS):      # plain runs such as "Slide No. "
                r.find("a:t", NS).text = ""
            for f in sh._element.findall(".//a:fld", NS):    # the live slide-number field
                f.find("a:t", NS).text = str(n)

def clear_body(slide):
    txBody = body_ph(slide)._element.find(".//p:txBody", NS)
    for p in txBody.findall("a:p", NS):
        txBody.remove(p)
    txBody.append(make_para("none", [("", False)], 10))

def caption(slide, x, y, w, h, text, size=13, color=NAVY, bold=False, align=PP_ALIGN.CENTER):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text; r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
    r.font.name = "Bookman Old Style"
    return tb

# ===================== 1. Title =====================
s = S[0]
boxes = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip() in ("Title", "Theme")]
tbox = [b for b in boxes if b.text_frame.text.strip() == "Title"][0]
thbox = [b for b in boxes if b.text_frame.text.strip() == "Theme"][0]
set_simple_text(tbox, "Connectome Lab: Running a Real Fly Brain as a Neural Network,\nwith an AI Agent that Discovers How It Works", size=24)
set_simple_text(thbox, "Theme: Artificial Intelligence, Neural Networks & Computational Neuroscience\nTeam: Govind  |  Bharath  |  Dhanush A  |  Shreenidhi Chalmi", size=20)
for b, sz in ((tbox, "2000"), (thbox, "1800")):
    ps = b._element.findall(".//a:p", NS)
    for r in ps[1].findall("a:r", NS):
        rp = r.find("a:rPr", NS); rp.set("sz", sz)
        sf = rp.find("a:solidFill", NS)
        if sf is not None: rp.remove(sf)
        sf = etree.Element("{%s}solidFill" % NS["a"]); etree.SubElement(sf, "{%s}srgbClr" % NS["a"], val="FFFFFF"); rp.insert(0, sf)

# ===================== 4. Introduction =====================
set_body(S[3], [
    (0, [("What it is about: ", True),
         ("In October 2024 the FlyWire project published every neuron and every connection in an adult fruit fly's "
          "brain — 139,255 neurons, 15 million connections — and released it free. In June 2026 two more maps followed "
          "(female brain + nerve cord, and the complete male nervous system). These are the first complete \"wiring "
          "diagrams\" of any animal that can walk, fly and learn.", False)]),
    (0, [("What we do with it: ", True),
         ("We run this real brain on a normal laptop as a neural network, connect it to a virtual fly so you can "
          "watch behaviour happen, and use it in two new ways: (1) as a neural-network architecture designed by "
          "evolution, and (2) as a lab where an AI agent runs its own experiments to find which neurons control "
          "which behaviour.", False)]),
    (0, [("Why it matters: ", True),
         ("A fly brain does vision, flight and navigation on ~1 milliwatt; a GPU doing similar work needs hundreds of "
          "watts. Flies share ~75% of human disease genes and are standard models for Parkinson's, Alzheimer's and "
          "epilepsy. Understanding how evolved wiring computes could give us cheaper AI and faster neuroscience.", False)]),
    (0, [("Already working: ", True),
         ("we loaded the real data and reproduced the headline result of Shiu et al. (Nature, Oct 2024) on a laptop — "
          "sugar on the tongue makes the feeding neuron fire (95–98 Hz every run); random input gives 0 Hz. "
          "Results and animation on slides 10–11.", False)]),
], size=18)

# ===================== 5. Literature Review =====================
set_body(S[4], [
    (0, [("Key studies:", True)]),
    (1, [("Dorkenwald et al. (2024, Nature) – ", True), ("the complete adult fly brain connectome (FlyWire): 139,255 neurons, 50 M synapses, with predicted transmitter for every connection.", False)]),
    (1, [("Shiu et al. (2024, Nature) – ", True), ("simulated the whole connectome with simple spiking neurons on a laptop; correctly predicted feeding, grooming and other behaviours, later confirmed in real flies.", False)]),
    (1, [("Lappalainen et al. (2024, Nature) – ", True), ("trained the fly visual-system wiring as a deep network in PyTorch; it predicted real neuron responses.", False)]),
    (1, [("Vaxenburg et al. (2025, Nature) – ", True), ("a physics-simulated fly body (MuJoCo) that neural controllers can drive — the basis for virtual-fly work.", False)]),
    (1, [("Costi et al. (2025, Biomimetics, ESA) – ", True), ("fly wiring as a reservoir computer; overfits far less than random networks of the same size.", False)]),
    (1, [("Jin et al. (2026, arXiv, FlyGM) – ", True), ("connectome as a graph neural network for reinforcement learning; beat standard MLP and GNN controllers.", False)]),
    (1, [("Bates et al. (2026, Nature) & MaleCNS (2026, Cell) – ", True), ("complete brain + nerve-cord connectomes, female and male; motor control turns out to be distributed in local loops.", False)]),
    (0, [("Trend: ", True), ("mapping (2020–24) → simulating (2024) → whole-body and male/female maps (2025–26) → using the wiring as an AI architecture (2025–26).", False)]),
    (0, [("Framework: ", True), ("spiking neural networks, graph theory, reservoir computing, connectome-constrained deep learning, LLM tool-using agents.", False)]),
    (0, [("Gap: ", True), ("no systematic test of evolved vs random/engineered wiring across many tasks; no autonomous AI experimentation on a connectome.", False)]),
], size=15, sub_size=14)

# ===================== 6. Problem Definition =====================
set_body(S[5], [
    (0, [("The problem: ", True),
         ("We now have the complete wiring of a brain, but wiring is not understanding. Finding out which neurons cause "
          "which behaviour still needs an expert to guess, stimulate, cut and repeat for months. And nobody knows "
          "whether a brain's evolved wiring is a better neural-network design than what humans build by hand.", False)]),
    (0, [("Context: ", True),
         ("The fly connectome was released in Oct 2024; brain + nerve cord and male versions in June 2026. Simulation "
          "code exists but is used by a handful of expert labs. Meanwhile AI data centres consume enormous power while "
          "a fly brain does complex tasks on almost none. The two fields have not been properly connected.", False)]),
    (0, [("Why it must be addressed: ", True),
         ("Without tools that turn structure into function automatically, the connectome stays a map few can read — "
          "and the mouse and human connectomes now in progress will hit the same bottleneck. Without a fair benchmark, "
          "AI research cannot tell whether biological wiring holds design lessons worth copying.", False)]),
    (0, [("Problem statement: ", True),
         ("Build a laptop-scale platform that runs the fruit-fly connectome as a neural network coupled to a virtual "
          "fly, lets an AI agent autonomously design and run stimulation/lesion experiments to identify behavioural "
          "circuits, and benchmarks the biological wiring against random and engineered networks on machine-learning "
          "tasks.", False)]),
], size=17)

# ===================== 7. Objectives =====================
set_body(S[6], [
    (0, [("Main objective: ", True),
         ("Turn the complete fruit-fly brain wiring into a runnable, trainable neural network on a laptop; use it to "
          "drive a virtual fly, to discover brain circuits automatically, and to test whether evolved wiring is a "
          "better AI architecture.", False)]),
    (0, [("Specific objectives:", True)]),
    (1, [("Brain engine: ", True), ("fast simulator of all 139,255 neurons with stimulate / record / cut; reproduce feeding, bitter-suppression and grooming from the 2024 Nature paper.", False)]),
    (1, [("Brain atlas: ", True), ("cell-type labels (sugar, bitter, touch, vision, motor) and graph queries such as \"shortest path from eye to legs\".", False)]),
    (1, [("Virtual fly (2-D): ", True), ("a fly in an on-screen arena driven by the brain's motor neurons — walks toward food, jumps from threats, grooms when touched; cut neurons live and watch behaviour break.", False)]),
    (1, [("AI scientist: ", True), ("an LLM with these tools plans, runs and interprets its own experiments and writes a report; validated by rediscovering known circuits.", False)]),
    (1, [("Neural-network benchmark: ", True), ("the wiring as a PyTorch model (frozen and trainable) vs random, small-world and MLP networks on 3–5 tasks, with a \"how much fly is left\" shuffle curve.", False)]),
    (1, [("Dashboard: ", True), ("a browser page where anyone can press \"sugar\", watch neurons light up, cut neurons, and chat with the AI scientist.", False)]),
    (0, [("Contribution: ", True), ("first autonomous-agent experimentation on a connectome; first systematic evolved-vs-random benchmark; an open tool usable without expert knowledge.", False)]),
], size=16, sub_size=15)

# ===================== 8. Methodology =====================
set_body(S[7], [
    (0, [("Design: ", True), ("quantitative and experimental — build the system, run controlled simulation experiments, compare with published lab data and with baseline networks.", False)]),
    (0, [("Data: ", True), ("public FlyWire v783 connectome (neurons, synapse counts, excitatory/inhibitory sign, cell-type annotations) from codex.flywire.ai. No data collection from people or animals — everything is simulation.", False)]),
    (0, [("Analysis: ", True), ("firing rates over 30 repeated trials; comparison with the Nature 2024 model; graph statistics (paths, hubs, communities); ML metrics (accuracy, data-efficiency, overfitting) with confidence intervals across seeds; behaviour metrics for the virtual fly (time to reach food, escape success).", False)]),
    (0, [("Tools: ", True), ("Python, NumPy/SciPy (sparse matrices), Pandas, NetworkX, PyTorch, Pygame/Canvas for the 2-D fly, Claude/OpenAI API for the agent, FastAPI + web front-end, GitHub. Runs on a laptop; free Colab GPU for the benchmark phase.", False)]),
    (0, [("Team roles: ", True), ("Brain engine & virtual fly (Govind) · Brain atlas & graph queries (Bharath) · AI scientist agent (Dhanush) · Neural-network benchmark & dashboard (Shreenidhi). One module each, all meeting in the dashboard.", False)]),
], size=16)

# ===================== 9. Methodology flowchart =====================
s = S[8]; clear_body(s)
def box(slide, x, y, w, h, title, sub, fill=NAVY):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = RGBColor(0x42, 0x71, 0x9B); shp.line.width = Pt(1.25); shp.shadow.inherit = False
    tf = shp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08); tf.margin_top = tf.margin_bottom = Inches(0.05)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title; r.font.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(255, 255, 255); r.font.name = "Bookman Old Style"
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = sub; r2.font.size = Pt(11); r2.font.color.rgb = RGBColor(255, 255, 255); r2.font.name = "Bookman Old Style"
def arrow(slide, x1, y1, x2, y2, head=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = RED; c.line.width = Pt(2.25)
    if head:
        etree.SubElement(c.line._get_or_add_ln(), "{%s}tailEnd" % NS["a"], type="triangle", w="med", len="med")
row1 = [("1. Load the brain", "FlyWire data: 139,255 neurons,\n15 M connections   ✔ done"),
        ("2. Build simulator", "spiking neurons in NumPy;\nstimulate / record / cut   ✔ done"),
        ("3. Validate", "reproduce feeding (✔ done),\nbitter, grooming vs Nature 2024"),
        ("4. Brain atlas", "cell-type labels + graph\nqueries (paths, hubs)")]
row2 = [("5. Virtual fly", "2-D fly driven by motor\nneurons: walk, escape, groom"),
        ("6. AI scientist", "LLM with tools plans and runs\nits own experiments"),
        ("7. NN benchmark", "fly wiring vs random / MLP\non learning tasks (PyTorch)"),
        ("8. Dashboard + paper", "web demo, report,\nconference submission")]
Wb, Hb, GAP, X0, Y1, Y2 = 2.75, 1.25, 0.42, 0.55, 1.55, 4.15
for i, (t, sub) in enumerate(row1):
    x = X0 + i * (Wb + GAP); box(s, x, Y1, Wb, Hb, t, sub)
    if i < 3: arrow(s, x + Wb, Y1 + Hb / 2, x + Wb + GAP, Y1 + Hb / 2)
for i, (t, sub) in enumerate(row2):
    x = X0 + i * (Wb + GAP); box(s, x, Y2, Wb, Hb, t, sub, fill=RED if i == 1 else NAVY)
    if i < 3: arrow(s, x + Wb, Y2 + Hb / 2, x + Wb + GAP, Y2 + Hb / 2)
x4, x5, ymid = X0 + 3 * (Wb + GAP) + Wb / 2, X0 + Wb / 2, (Y1 + Hb + Y2) / 2
arrow(s, x4, Y1 + Hb, x4, ymid); arrow(s, x4, ymid, x5, ymid, head=False); arrow(s, x5, ymid, x5, Y2)
caption(s, 0.55, 5.65, 12.2, 0.9, "Steps 1–2 and the first validation are done (Sept 2026). Each later step is owned by one team "
        "member and plugs into the dashboard. Every experiment is repeated over 30 trials and compared with published results.", size=14)

# ===================== 10. Preliminary results (new slide) =====================
s = S[9]; set_simple_text(title_box(s), "Preliminary Results"); clear_body(s); set_slide_number(s, 10)
s.shapes.add_picture("../fly/sugar_test.png", Inches(0.45), Inches(1.35), width=Inches(6.6))
s.shapes.add_picture("../fly/lesion_test.png", Inches(7.35), Inches(1.35), width=Inches(5.6))
caption(s, 0.45, 4.05, 6.6, 0.6, "Sugar on the tongue → 380 neurons active, feeding neuron MN9 at 95–98 Hz.\nRandom input → MN9 silent (0 Hz). Same result in 3 of 3 runs.", size=12)
caption(s, 7.35, 4.05, 5.6, 0.6, "Brain surgery: cut the 3 command neurons → feeding stops (1 Hz).\nCut 10 random neurons → nothing changes (91 Hz).", size=11)
set_body(s, [
    (0, [("What this proves: ", True), ("the real 139,255-neuron brain runs on a normal laptop (1 s of brain time ≈ 35 s), matches the Nature 2024 numbers (~400 active neurons, MN9 firing on sugar), and we can already find and cut a circuit — the tongue-to-mouth path is just 3 synapses long.", False)]),
], size=15)
body_ph(s).top = Inches(4.85); body_ph(s).height = Inches(1.9)

# ===================== 11. Live demo (new slide) =====================
s = S[10]; set_simple_text(title_box(s), "Live Demo"); clear_body(s); set_slide_number(s, 11)
pic = s.shapes.add_picture("../fly/fly_feeding_demo.gif", Inches(0.9), Inches(1.15), width=Inches(9.3))
set_body(s, [
    (0, [("Animation (plays in slideshow): ", True), ("left — the real wiring, neurons placed by distance from the tongue, flashing as they fire. Right — a fly whose mouth is driven live by the MN9 'eat' neuron.", False)]),
    (0, [("0–350 ms: ", True), ("sugar arrives → signal ripples hop by hop → MN9 fires → mouth extends.", False)]),
    (0, [("350 ms: ", True), ("we cut 3 command neurons (red ✕) → MN9 goes silent → mouth retracts.", False)]),
    (0, [("December version: ", True), ("a 2-D fly walking in an arena, buttons for sugar / bitter / touch / light, a scalpel tool, and a chat box to the AI scientist.", False)]),
], size=13, sub_size=12)
b = body_ph(s); b.left = Inches(10.3); b.top = Inches(1.15); b.width = Inches(2.85); b.height = Inches(5.5)

# ===================== 12. Proposed Outcomes =====================
set_body(S[11], [
    (0, [("Expected results:", True)]),
    (1, "A whole-brain simulator on a laptop that matches the Nature 2024 results for feeding, bitter suppression and grooming."),
    (1, "A 2-D virtual fly driven by the brain: walks toward food, escapes threats, grooms — and visibly breaks when specific neurons are cut."),
    (1, "An AI scientist that, given only \"what makes the fly eat?\", finds the responsible neurons by itself (we already found the 3 key neurons by hand)."),
    (1, "A benchmark table and shuffle curve showing where fly wiring beats, ties or loses to random and engineered networks."),
    (1, "An open-source web dashboard anyone can use without neuroscience knowledge."),
    (0, [("Impact: ", True), ("labs simulate before experimenting (saves months); AI researchers get evidence on copying biological wiring; the pipeline transfers to the mouse connectome now being mapped; supports low-power (neuromorphic) computing.", False)]),
    (0, [("Future directions: ", True), ("3-D physics fly (NeuroMechFly / Flybody); learning via dopamine-gated plasticity in the mushroom body; male-vs-female connectome comparison; running the circuit on a neuromorphic chip or Raspberry Pi.", False)]),
], size=16, sub_size=15)

# ===================== 13. Timeline =====================
set_body(S[12], [
    (0, [("Phase 1 – Foundation (Weeks 1–2, 17 Sep – 1 Oct)", True)]),
    (1, "Scope agreed with mentor; GitHub repo; simulator cleaned up (30-trial runs, faster updates); cell-type labels loaded."),
    (0, [("Phase 2 – Validation + Atlas (Weeks 3–5, 1 Oct – 22 Oct)", True)]),
    (1, "Reproduce feeding, bitter-suppression, grooming; \"we match the paper\" table; graph queries (paths, hubs, communities)."),
    (0, [("Phase 3 – Virtual fly + AI scientist (Weeks 6–9, 22 Oct – 19 Nov)", True)]),
    (1, "2-D fly arena: motor neurons → movement; food-seeking and escape behaviours; live neuron cutting."),
    (1, "Expose stimulate / record / cut / query as LLM tools; agent must rediscover the feeding circuit unaided, then grooming."),
    (0, [("Phase 4 – Neural-network benchmark (Weeks 8–11, 5 Nov – 3 Dec)", True)]),
    (1, "Fly wiring as frozen (reservoir) and trainable PyTorch models; 3–5 tasks; shuffle curve; baselines."),
    (0, [("Phase 5 – Dashboard, report, paper (Weeks 11–13, 3 Dec – 17 Dec)", True)]),
    (1, "Web dashboard integrating fly, brain view and chat; final demo; report; IEEE-format paper draft."),
], size=16, sub_size=15)

# ===================== 14. Conclusion =====================
set_body(S[13], [
    (0, [("Summary: ", True), ("For the first time we have a complete brain wiring diagram, and it runs on a laptop — we have already reproduced a Nature 2024 result and animated it. Our project turns this into a platform: a simulator, a labelled atlas, a virtual fly, an AI agent that experiments on the brain by itself, and a fair test of whether evolution's wiring beats human-designed networks.", False)]),
    (0, [("Final thought: ", True), ("This is where neuroscience and AI meet. A fly solves vision, flight and learning with 139 thousand neurons and almost no power. If its wiring holds design lessons, the fastest way to find them is to run it, poke it, and let an AI help us understand it.", False)]),
    (0, [("Next steps: ", True), ("mentor approval → validation runs (Oct) → virtual fly and agent prototype (Nov) → benchmark and dashboard (Dec) → conference paper.", False)]),
], size=18)

# ===================== 15–16. References =====================
REFS_A = [
    "Dorkenwald, S., Matsliah, A., Sterling, A. R., Schlegel, P., Yu, S.-C., McKellar, C. E., … Murthy, M. (2024). Neuronal wiring diagram of an adult brain. Nature, 634, 124–138. https://doi.org/10.1038/s41586-024-07558-y",
    "Schlegel, P., Yin, Y., Bates, A. S., Dorkenwald, S., Eichler, K., … Jefferis, G. S. X. E. (2024). Whole-brain annotation and multi-connectome cell typing of Drosophila. Nature, 634, 139–152. https://doi.org/10.1038/s41586-024-07686-5",
    "Lin, A., Yang, R., Dorkenwald, S., Matsliah, A., Sterling, A. R., … Murthy, M. (2024). Network statistics of the whole-brain connectome of Drosophila. Nature, 634, 153–165. https://doi.org/10.1038/s41586-024-07968-y",
    "Shiu, P. K., Sterne, G. R., Spiller, N., Blagburn, D. I., Mosca, T. J., … Scott, K. (2024). A Drosophila computational brain model reveals sensorimotor processing. Nature, 634, 210–219. https://doi.org/10.1038/s41586-024-07763-9",
    "Lappalainen, J. K., Tschopp, F. D., Prakhya, S., McGill, M., Nern, A., … Turaga, S. C. (2024). Connectome-constrained networks predict neural activity across the fly visual system. Nature, 634, 1132–1140. https://doi.org/10.1038/s41586-024-07939-3",
    "Nern, A., Loesche, F., Takemura, S., Burnett, L. E., Dreher, M., … Reiser, M. B. (2025). Connectome-driven neural inventory of a complete visual system. Nature, 641, 1225–1237. https://doi.org/10.1038/s41586-025-08746-0",
    "Vaxenburg, R., Siwanowicz, I., Merel, J., Robie, A. A., Morrow, C., … Turaga, S. C. (2025). Whole-body physics simulation of fruit fly locomotion. Nature, 643, 1312–1320. https://doi.org/10.1038/s41586-025-09029-4",
    "Stürner, T., et al. (2025). Comparative connectomics of Drosophila descending and ascending neurons. Nature, 643, 158–172. https://doi.org/10.1038/s41586-025-08925-z",
]
REFS_B = [
    "Costi, L., Hadjiivanov, A., Dold, D., Hale, Z. F., & Izzo, D. (2025). The Drosophila connectome as a computational reservoir for time-series prediction. Biomimetics, 10(5), 341. https://doi.org/10.3390/biomimetics10050341",
    "Bates, A. S., Phelps, J. S., Kim, M., Yang, H. H., … BANC-FlyWire Consortium. (2026). Distributed control circuits across a brain-and-cord connectome. Nature. https://doi.org/10.1038/s41586-026-10735-w",
    "FlyEM Project Team (HHMI Janelia, University of Cambridge, MRC LMB & Google Research). (2026). Sexual dimorphism in the complete connectome of the Drosophila male central nervous system. Cell. https://www.cell.com/cell/fulltext/S0092-8674(26)00942-6",
    "Jin, Z., Zhu, Y., Zhang, C., & Sui, Y. (2026). Whole-brain connectomic graph model enables whole-body locomotion control in fruit fly. arXiv preprint arXiv:2602.17997.",
    "Lobato-Rios, V., et al. (2022). NeuroMechFly, a neuromechanical model of adult Drosophila melanogaster. Nature Methods, 19, 620–627. https://doi.org/10.1038/s41592-022-01466-7",
    "State of Brain Emulation Report 2025. (2025). arXiv preprint arXiv:2510.15745.",
    "FlyWire Consortium. (2024). FlyWire Codex: Connectome Data Explorer, release v783 [Dataset]. https://codex.flywire.ai",
    "Shiu, P. K. (2024). Drosophila_brain_model [Source code]. GitHub. https://github.com/philshiu/Drosophila_brain_model",
    "Janelia FlyEM. (2026). Male CNS Connectome v1.0 [Dataset]. https://male-cns.janelia.org",
]
set_body(S[14], [("none", r) for r in REFS_A], size=12, sub_size=12); set_simple_text(title_box(S[14]), "References (1/2)")
set_body(S[15], [("none", r) for r in REFS_B], size=12, sub_size=12); set_simple_text(title_box(S[15]), "References (2/2)")

# ===================== 17. Expected outcome =====================
s = S[16]
set_body(s, [
    (0, [("Our target outcome: ", True), ("a conference publication plus an open-source software release.", False)]),
    (0, [("Conference publication: ", True), ("an IEEE-format paper on \"autonomous LLM-driven circuit discovery in a whole-brain connectome\" and the \"evolved-vs-random wiring\" benchmark. Target venues: IEEE conferences on computational intelligence / neural networks, or a computational-neuroscience workshop (COSYNE / NeurIPS workshops).", False)]),
    (0, [("Open-source release: ", True), ("the simulator, atlas, virtual fly, agent tools and dashboard on GitHub with documentation, so other students and labs can reproduce and extend our work.", False)]),
    (0, [("Research proposal (stretch): ", True), ("extend the pipeline to the 3-D physics fly, the mouse connectome and neuromorphic hardware, with a budget for a small GPU workstation and a Loihi / Raspberry Pi test board.", False)]),
    (0, [("Patent: ", True), ("not planned — the value here is in open science and publication.", False)]),
], size=17)
for sh in s.shapes:
    if sh.is_placeholder and sh.placeholder_format.type == 16:
        for t in sh._element.findall(".//a:t", NS):
            t.text = "Friday, 18 September 2026"

# slide numbers for every slide after the insertions (fields are cached text; keep them right)
for i, sl in enumerate(S, 1):
    if i >= 10:
        set_slide_number(sl, i)

prs.save(OUT); print("saved", OUT)
